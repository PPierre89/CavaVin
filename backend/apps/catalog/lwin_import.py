"""Ingestion du dump du référentiel LWIN (Liv-ex) dans ``ReferenceLwin``.

Logique partagée entre la commande ``manage.py import_lwin`` (usage CLI /
conteneur) et l'endpoint d'upload du panneau d'administration. Accepte le XLSX
Liv-ex d'origine ou un export CSV équivalent. Idempotent : les entrées
existantes (même code LWIN) sont mises à jour, les nouvelles créées.
"""

from __future__ import annotations

import csv

from .enrichment.lwin import vider_cache_referentiel
from .enrichment.normalize import guess_couleur
from .models import ReferenceLwin

# Taille des lots bulk_create : compromis mémoire / nombre de requêtes SQLite.
_LOT = 1000

_CHAMPS_MAJ = [
    "producteur", "vin", "pays", "region", "sous_region", "couleur", "classification",
]


class LwinImportError(Exception):
    """Fichier LWIN illisible / introuvable (message destiné à l'utilisateur)."""


def _texte(valeur) -> str:
    """Valeur de cellule -> chaîne nettoyée.

    Le dump encode l'absence par la chaîne « NA » (ramenée à vide), et la
    lecture XLSX renvoie les codes LWIN en nombres (1000001.0 -> "1000001")."""
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    texte = str(valeur).strip()
    return "" if texte.upper() == "NA" else texte


def _lignes(chemin: str, delimiter: str):
    """Itère les lignes du dump en dictionnaires {EN-TÊTE: valeur brute},
    que le fichier soit le XLSX Liv-ex d'origine ou un export CSV."""
    if chemin.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - dépendance épinglée
            raise LwinImportError("openpyxl est requis pour lire un dump XLSX.") from exc
        try:
            classeur = openpyxl.load_workbook(chemin, read_only=True)
        except (OSError, ValueError) as exc:
            raise LwinImportError(f"Impossible d'ouvrir {chemin} : {exc}") from exc
        feuille = classeur[classeur.sheetnames[0]]
        lignes = feuille.iter_rows(values_only=True)
        entetes = [(_texte(c)).upper() for c in next(lignes, [])]
        for valeurs in lignes:
            yield dict(zip(entetes, valeurs))
        classeur.close()
    else:
        try:
            fichier = open(chemin, newline="", encoding="utf-8-sig")
        except OSError as exc:
            raise LwinImportError(f"Impossible d'ouvrir {chemin} : {exc}") from exc
        with fichier:
            for brut in csv.DictReader(fichier, delimiter=delimiter):
                yield {(k or "").strip().upper(): v for k, v in brut.items()}


def _enregistrer(lot: list[ReferenceLwin]) -> int:
    """Upsert d'un lot : création, ou mise à jour si le code LWIN existe déjà."""
    if not lot:
        return 0
    ReferenceLwin.objects.bulk_create(
        lot,
        update_conflicts=True,
        unique_fields=["lwin"],
        update_fields=_CHAMPS_MAJ,
    )
    return len(lot)


def importer_lwin(chemin: str, delimiter: str = ",") -> int:
    """Ingest le dump LWIN situé à ``chemin``. Retourne le nombre d'entrées
    importées/mises à jour. Lève ``LwinImportError`` si le fichier est illisible."""
    total = 0
    lot: list[ReferenceLwin] = []
    for ligne in _lignes(chemin, delimiter):
        code = _texte(ligne.get("LWIN"))
        if not code:
            continue
        # Seules les entrées actives du référentiel sont importées.
        statut = _texte(ligne.get("STATUS")).lower()
        if statut and statut != "live":
            continue
        producteur = " ".join(
            p for p in (_texte(ligne.get("PRODUCER_TITLE")), _texte(ligne.get("PRODUCER_NAME"))) if p
        ) or _texte(ligne.get("DISPLAY_NAME"))
        if not producteur:
            continue
        lot.append(
            ReferenceLwin(
                lwin=code[:16],
                producteur=producteur[:255],
                vin=_texte(ligne.get("WINE"))[:255],
                pays=_texte(ligne.get("COUNTRY"))[:100],
                region=_texte(ligne.get("REGION"))[:255],
                sous_region=_texte(ligne.get("SUB_REGION"))[:255],
                # L'effervescence est portée par SUB_TYPE (« Sparkling » ;
                # TYPE vaut toujours « Wine ») et la couleur par COLOUR.
                couleur=guess_couleur(
                    _texte(ligne.get("SUB_TYPE")), _texte(ligne.get("COLOUR"))
                ),
                classification=_texte(ligne.get("CLASSIFICATION"))[:255],
            )
        )
        if len(lot) >= _LOT:
            total += _enregistrer(lot)
            lot = []
    total += _enregistrer(lot)
    # bulk_create ne déclenche pas les signaux du modèle : on invalide à la main
    # l'index en mémoire pour que le nouveau dump soit servi immédiatement.
    vider_cache_referentiel()
    return total
